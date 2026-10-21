"""Gera um laudo de exames sintético em DOCX, XLSX e MD (dados/).

Seis exames, paciente fictício. O laboratório marca em negrito o que está fora da
referência (DOCX e MD); a XLSX tem a mesma tabela, sem formatação.
"""
from pathlib import Path

from docx import Document
from openpyxl import Workbook

PASTA = Path(__file__).parent / "dados"
TITULO = "Laboratório Fictício Ltda - Laudo de exames"
PACIENTE = "Paciente: Maria Teste da Silva (dado sintético) - Coleta: 01/10/2026"
LEGENDA = "Valores em negrito estão fora da referência."
CABECALHO = ["Exame", "Resultado", "Unidade", "Referência"]
# (exame, resultado, unidade, referência, negrito no laudo)
EXAMES = [
    ("Hemoglobina", "11,2", "g/dL", "12,0 a 15,5", True),
    ("Leucócitos", "7.850", "/mm³", "4.000 a 11.000", False),
    ("Glicose em jejum", "126", "mg/dL", "70 a 99", True),
    ("Creatinina", "0,9", "mg/dL", "0,5 a 1,1", False),
    ("Potássio", "5,8", "mEq/L", "3,5 a 5,1", True),
    ("TSH", "2,1", "µUI/mL", "0,4 a 4,0", False),
]


def gerar_docx(caminho: Path) -> None:
    doc = Document()
    doc.add_heading(TITULO, level=1)
    doc.add_paragraph(PACIENTE)
    tabela = doc.add_table(rows=1, cols=len(CABECALHO))
    for celula, texto in zip(tabela.rows[0].cells, CABECALHO):
        celula.text = texto
    for exame, resultado, unidade, referencia, negrito in EXAMES:
        celulas = tabela.add_row().cells
        celulas[0].text, celulas[2].text, celulas[3].text = exame, unidade, referencia
        run = celulas[1].paragraphs[0].add_run(resultado)
        run.bold = negrito
    doc.add_paragraph(LEGENDA)
    doc.save(caminho)


def gerar_xlsx(caminho: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(CABECALHO)
    for exame, resultado, unidade, referencia, _ in EXAMES:
        ws.append([exame, resultado, unidade, referencia])
    wb.save(caminho)


def gerar_md(caminho: Path) -> None:
    linhas = [f"# {TITULO}", "", PACIENTE, "", "| " + " | ".join(CABECALHO) + " |",
              "|" + "---|" * len(CABECALHO)]
    for exame, resultado, unidade, referencia, negrito in EXAMES:
        valor = f"**{resultado}**" if negrito else resultado
        linhas.append(f"| {exame} | {valor} | {unidade} | {referencia} |")
    linhas += ["", LEGENDA, ""]
    caminho.write_text("\n".join(linhas), encoding="utf-8")


def gerar() -> list[Path]:
    PASTA.mkdir(exist_ok=True)
    saidas = [PASTA / "laudo_sintetico.docx", PASTA / "laudo_sintetico.xlsx",
              PASTA / "laudo_sintetico.md"]
    gerar_docx(saidas[0])
    gerar_xlsx(saidas[1])
    gerar_md(saidas[2])
    return saidas


if __name__ == "__main__":
    for caminho in gerar():
        print("gerado:", caminho.relative_to(Path(__file__).parent))
